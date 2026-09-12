import os
import sys
import json
import time
import glob
import sqlite3
import argparse
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = Path(__file__).resolve().parent
REPO_DIR = BASE_DIR / 'Repo'
REPO_DIR.mkdir(parents=True, exist_ok=True)

DRIVEFS_ROOT = Path(os.environ.get('LOCALAPPDATA', '')) / 'Google' / 'DriveFS'

def find_drivefs_db() -> str | None:
    if not DRIVEFS_ROOT.exists():
        return None
    dbs = glob.glob(str(DRIVEFS_ROOT / '**' / 'mirror_metadata_sqlite.db'), recursive=True)
    return dbs[0] if dbs else None

def get_drive_file_id(filename: str, max_wait_sec: int = 15) -> str | None:
    db_path = find_drivefs_db()
    if not db_path:
        return None

    # Esperar a que DriveFS sincronice y registre el archivo en la base local
    for _ in range(max_wait_sec // 2):
        time.sleep(2)
        try:
            conn = sqlite3.connect(f'file:{db_path}?mode=ro', uri=True)
            cur = conn.cursor()
            cur.execute('SELECT id FROM items WHERE local_title = ? AND trashed = 0', (filename,))
            row = cur.fetchone()
            conn.close()
            if row and row[0]:
                return row[0]
        except Exception:
            pass
    return None

def save_document(title: str, content: str) -> dict:
    # 1. Crear documento Word .docx limpio y legible
    doc = docx.Document()
    
    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = content.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue
            
        # Detectar encabezado o separador
        if stripped.startswith('---'):
            p = doc.add_paragraph()
            p_run = p.add_run('-' * 40)
            p_run.font.color.rgb = RGBColor(180, 180, 180)
        elif any(stripped.startswith(f"{i})") for i in range(1, 8)) or stripped.startswith("FUENTES"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(26, 82, 118) # Azul corporativo
        elif stripped == title.upper():
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(16)
        elif stripped.startswith("Fecha:"):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(100, 100, 100)
        else:
            p = doc.add_paragraph(stripped)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)

    docx_filename = f"{title}.docx"
    docx_path = REPO_DIR / docx_filename
    doc.save(docx_path)

    # 2. Guardar también versión .txt en texto plano exacto
    txt_filename = f"{title}.txt"
    txt_path = REPO_DIR / txt_filename
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write(content)

    # 3. Obtener URL de Drive buscando el ID sincronizado
    file_id = get_drive_file_id(docx_filename)
    if not file_id:
        file_id = get_drive_file_id(txt_filename)

    if file_id:
        url = f"https://docs.google.com/document/d/{file_id}/edit"
    else:
        # Fallback con ID de carpeta Drive
        url = f"https://drive.google.com/drive/folders/1fi9RKQYFVAPUjN-0vChYYtrlFWQGxkcC"

    return {
        'success': True,
        'id': file_id,
        'title': title,
        'url': url,
        'local_docx': str(docx_path),
        'local_txt': str(txt_path)
    }

def main():
    parser = argparse.ArgumentParser(description="Guardar y sincronizar artículo en Google Drive Repo")
    parser.add_argument("--title", required=True, help="Título del documento")
    parser.add_argument("--content", help="Contenido en texto plano")
    parser.add_argument("--file", help="Ruta al archivo con contenido")
    args = parser.parse_args()

    content = args.content
    if args.file:
        with open(args.file, 'r', encoding='utf-8') as f:
            content = f.read()

    if not content:
        print("Error: Se requiere --content o --file")
        sys.exit(1)

    res = save_document(args.title, content)
    print(json.dumps(res, indent=2, ensure_ascii=False))

if __name__ == '__main__':
    main()
